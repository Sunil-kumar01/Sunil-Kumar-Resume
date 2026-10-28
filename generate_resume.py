import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, KeepTogether
from xml.sax.saxutils import escape

parser = argparse.ArgumentParser()
parser.add_argument("--data", default="resume_content.json")
parser.add_argument("--prefix", default="Sunil_Kumar_Resume")
args = parser.parse_args()

ROOT = Path(__file__).parent
DATA = json.loads((ROOT / args.data).read_text())
NAVY = RGBColor(18, 53, 91)
BLUE = RGBColor(40, 101, 163)

# Word version
word = Document()
section = word.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.7)
section.top_margin = section.bottom_margin = Inches(0.52)
section.left_margin = section.right_margin = Inches(0.62)
normal = word.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(9)
normal.font.color.rgb = RGBColor(35, 48, 68)
normal.paragraph_format.space_after = Pt(2)
normal.paragraph_format.line_spacing = 1.0

p = word.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(DATA["name"]); r.bold = True; r.font.name = "Arial"; r.font.size = Pt(24); r.font.color.rgb = NAVY
p = word.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(DATA["title"]); r.bold = True; r.font.name = "Arial"; r.font.size = Pt(9.5); r.font.color.rgb = BLUE
p = word.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run(f"Phone: {DATA['phone']}  |  Email: {DATA['email']}  |  GitHub: {DATA['github']}")

def heading(text):
    p = word.add_paragraph(); p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text.upper()); r.bold = True; r.font.name = "Arial"; r.font.size = Pt(10.5); r.font.color.rgb = NAVY

def bullet(text):
    p = word.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.18); p.paragraph_format.first_line_indent = Inches(-0.12); p.paragraph_format.space_after = Pt(1)
    p.add_run(text)

heading("Professional Summary")
word.add_paragraph(DATA["summary"])
heading("Core Skills")
for label, text in DATA["skills"]:
    bullet(f"{label}: {text}")
heading("Professional Experience")
for job in DATA["experience"]:
    p = word.add_paragraph(); p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(1); p.paragraph_format.keep_with_next = True
    r = p.add_run(job["role"] + "  "); r.bold = True; r.font.color.rgb = BLUE
    r = p.add_run(job["dates"]); r.font.size = Pt(8.5)
    if job.get("subrole"):
        p = word.add_paragraph(job["subrole"]); p.paragraph_format.space_after = Pt(1)
        for r in p.runs: r.italic = True; r.font.size = Pt(8.3)
    for item in job["bullets"]: bullet(item)
heading("Personal AI Projects")
for project in DATA["projects"]:
    p = word.add_paragraph(); p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(1)
    r = p.add_run(project["title"]); r.bold = True; r.font.color.rgb = BLUE
    p = word.add_paragraph(f"{project['description']} Stack: {project['stack']}. Repository: {project['url']}")
    p.paragraph_format.space_after = Pt(3)
heading("Education & Professional Development")
for item in DATA["education"]: bullet(item)
word.save(ROOT / f"{args.prefix}.docx")

# PDF version
styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name="NameCustom", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=23, leading=26, textColor=colors.HexColor("#12355b"), alignment=TA_CENTER, spaceAfter=2))
styles.add(ParagraphStyle(name="TagCustom", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=9, leading=11, textColor=colors.HexColor("#2865a3"), alignment=TA_CENTER, spaceAfter=5))
styles.add(ParagraphStyle(name="ContactCustom", parent=styles["Normal"], fontSize=8, leading=10, alignment=TA_CENTER, spaceAfter=8))
styles.add(ParagraphStyle(name="SectionCustom", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, leading=12, textColor=colors.HexColor("#12355b"), spaceBefore=7, spaceAfter=3, keepWithNext=True))
styles.add(ParagraphStyle(name="BodyCustom", parent=styles["BodyText"], fontSize=8.2, leading=10, spaceAfter=3))
styles.add(ParagraphStyle(name="RoleCustom", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=8.6, leading=10, textColor=colors.HexColor("#1c4e7a"), spaceBefore=4, spaceAfter=1, keepWithNext=True))
styles.add(ParagraphStyle(name="BulletCustom", parent=styles["BodyText"], fontSize=7.8, leading=9.2, leftIndent=11, firstLineIndent=-7, spaceAfter=1))
flow = [Paragraph(escape(DATA["name"]), styles["NameCustom"]), Paragraph(escape(DATA["title"]), styles["TagCustom"]), Paragraph(f"Phone: {escape(DATA['phone'])} &nbsp; | &nbsp; Email: {escape(DATA['email'])} &nbsp; | &nbsp; GitHub: {escape(DATA['github'])}", styles["ContactCustom"])]
def pdf_heading(text): flow.append(Paragraph(escape(text.upper()), styles["SectionCustom"]))
def pdf_bullet(text): flow.append(Paragraph("• &nbsp; " + escape(text), styles["BulletCustom"]))
pdf_heading("Professional Summary")
flow.append(Paragraph(escape(DATA["summary"]), styles["BodyCustom"]))
pdf_heading("Core Skills")
rows = []
for i in range(0, len(DATA["skills"]), 2):
    row = []
    for label, text in DATA["skills"][i:i+2]: row.append(Paragraph(f"<b>{escape(label)}:</b> {escape(text)}", styles["BodyCustom"]))
    if len(row) == 1: row.append("")
    rows.append(row)
t = Table(rows, colWidths=[3.45*inch, 3.45*inch], hAlign="LEFT")
t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#f2f6fa")),("BOX",(0,0),(-1,-1),0.4,colors.HexColor("#d5e0ea")),("INNERGRID",(0,0),(-1,-1),0.3,colors.HexColor("#d5e0ea")),("VALIGN",(0,0),(-1,-1),"TOP"), ("LEFTPADDING",(0,0),(-1,-1),5),("RIGHTPADDING",(0,0),(-1,-1),5),("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),2)]))
flow.append(t)
pdf_heading("Professional Experience")
for job in DATA["experience"]:
    elements = [Paragraph(f"{escape(job['role'])} &nbsp; <font color='#59697a'>{escape(job['dates'])}</font>", styles["RoleCustom"])]
    if job.get("subrole"): elements.append(Paragraph(f"<i>{escape(job['subrole'])}</i>", styles["BodyCustom"]))
    elements.extend(Paragraph("• &nbsp; " + escape(item), styles["BulletCustom"]) for item in job["bullets"])
    flow.append(KeepTogether(elements))
pdf_heading("Personal AI Projects")
for project in DATA["projects"]:
    content = f"<b><font color='#1c4e7a'>{escape(project['title'])}</font></b><br/>{escape(project['description'])} <b>Stack:</b> {escape(project['stack'])}. <link href='{escape(project['url'])}' color='#175f9c'>{escape(project['link_label'])}</link>."
    box = Table([[Paragraph(content, styles["BodyCustom"])]], colWidths=[6.9*inch])
    box.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#f2f6fa")),("LINEBEFORE",(0,0),(0,0),2,colors.HexColor("#3c78a8")),("LEFTPADDING",(0,0),(-1,-1),7),("RIGHTPADDING",(0,0),(-1,-1),6),("TOPPADDING",(0,0),(-1,-1),5),("BOTTOMPADDING",(0,0),(-1,-1),3)]))
    flow.extend([box, Spacer(1,3)])
pdf_heading("Education & Professional Development")
for item in DATA["education"]: pdf_bullet(item)
SimpleDocTemplate(str(ROOT / f"{args.prefix}.pdf"), pagesize=A4, rightMargin=.62*inch, leftMargin=.62*inch, topMargin=.48*inch, bottomMargin=.48*inch, title=f"{DATA['name']} Resume", author=DATA["name"]).build(flow)
